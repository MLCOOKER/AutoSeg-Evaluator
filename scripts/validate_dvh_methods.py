"""Score candidate DVH methods against DVHs whose true values are known.

Roadmap item #7. AutoSeg computes structure-set DVHs today with dicompyler-core,
which samples each contour on the dose grid. The alternative is to compute them
from the 3D masks the geometric metrics already use. This script scores every
candidate against two sets of analytic ground truth and writes the tables the
decision is made from.

Benchmarks
----------
Nelms et al. 2015 (Med Phys 42:4435, doi:10.1118/1.4923175)
    A sphere, cylinders and cones (3.6-12 cc) contoured every 0.2-3 mm, in
    1 Gy/mm linear dose fields on 0.4-3 mm dose grids. The truth is the solid
    itself, extended half a slice at each end, so a method is also charged for
    how it fills the gap between contour planes. Tests 1-3 are reproduced as
    published, beside the paper's own results for Pinnacle3 and PlanIQ.

Disc phantoms (built here)
    Stacks of circular contours (sphere, cylinder, ring with a hole), radius
    2.5-20 mm, at random sub-voxel positions on two clinical grids, in linear
    dose fields along x, along z and at 45 degrees between them. Each contour
    stands for a slab one slice thick, and every slab has a closed-form DVH, so
    the truth is exact for the contours as drawn. This isolates each method's
    own error from the question of how contours are joined.

Large structures (built here)
    One sphere or cylinder per size from 4 cc to 6,300 cc on the tender
    cohort's grid, to time each method against its accuracy where the number
    of samples, not the outline, sets the cost.

Methods
-------
dicompyler     Production today (``core.dvh.compute_dvh_metrics``): dicompyler-core
               tests each dose-grid point in each contour plane.
dicompyler-ss  The same, supersampled in-plane to a quarter of the dose pixel on
               every structure (production does this only as a zero-volume retry).
mask           The 3D mask the geometric metrics use (shared reading, half-open
               fill: ``core.masks.mask_with_reading``), with the dose interpolated
               trilinearly at each voxel centre.
mask-ss        The same voxels, with the dose integrated over each voxel by
               sub-samples no further apart than ``--subsample-mm``.
polygon        No voxels: the shared reading's regions themselves, weighted by the
               exact area they cover of every sub-cell, with the same dose
               sub-sampling. This is a DVH from the raw polygons.

Usage::

    python scripts/validate_dvh_methods.py --nelms <Nelms dataset folder> \\
        --out docs/DVH_METHOD_VALIDATION.md [--csv results.csv]

Every input is synthetic; nothing here comes from a patient.
"""

from __future__ import annotations

import argparse
import csv
import datetime
import importlib.metadata
import math
import os
import platform
import subprocess
import sys
import tempfile
import time
import warnings
from collections import defaultdict
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field, replace
from pathlib import Path

import numpy as np
import pydicom
import shapely
import SimpleITK as sitk
from pydicom.dataset import Dataset, FileDataset
from pydicom.uid import ExplicitVRLittleEndian, RTDoseStorage, RTStructureSetStorage, generate_uid
from scipy import ndimage, optimize
from shapely.geometry.polygon import orient

from autoseg_evaluator import __version__ as autoseg_version
from autoseg_evaluator.core.dvh import (
    DVHConfig,
    _find_roi_contour,
    _single_plane_thickness,
    _statistic_value,
    _supersample_resolution,
    compute_dvh_metrics,
)
from autoseg_evaluator.core.masks import _fill_structure, mask_with_reading

warnings.filterwarnings("ignore", module="pydicom")

METHODS = ("dicompyler", "dicompyler-ss", "mask", "mask-ss", "polygon")
SAMPLED = ("mask-ss", "polygon")  # the methods that take a sub-sample spacing
METRICS = ("volume_cc", "dmin", "dmax", "dmean", "d99", "d95", "d5", "d1", "d0.03cc")
LABEL = {
    "volume_cc": "V",
    "dmin": "Dmin",
    "dmax": "Dmax",
    "dmean": "Dmean",
    "d99": "D99",
    "d95": "D95",
    "d5": "D5",
    "d1": "D1",
    "d0.03cc": "D0.03cc",
}
# Nelms et al. set Dmin and Dmax aside as the least clinically relevant; so do we
# wherever a summary says "clinical".
CLINICAL = ("dmean", "d99", "d95", "d5", "d1", "d0.03cc")
CONFIG = DVHConfig(
    include_dmin=True,
    include_dmean=True,
    include_dmax=True,
    d_at_volumes_pct=[99.0, 95.0, 5.0, 1.0],
    d_at_volumes_cc=[0.03],
    v_at_doses_gy=[],
)
PRODUCTION_KEY = {
    "dmin_gy": "dmin",
    "dmax_gy": "dmax",
    "dmean_gy": "dmean",
    "d99_gy": "d99",
    "d95_gy": "d95",
    "d5_gy": "d5",
    "d1_gy": "d1",
    "d0.03cc_gy": "d0.03cc",
}
CHUNK = 2_000_000  # points per interpolation call, to bound memory
BIN_GY = 0.001  # dose histogram bin of the sampled methods
# A sub-cell covered by less than this share is floating-point residue from an
# edge lying along a grid line, and is dropped.
SLIVER = 1e-9


# --------------------------------------------------------------------------
# Dose
# --------------------------------------------------------------------------


@dataclass
class DoseGrid:
    """An RT Dose grid in Gy, sampled by trilinear interpolation."""

    values: np.ndarray  # (frames, rows, columns)
    origin_xy: np.ndarray  # x, y of the first voxel centre, mm
    spacing_xy: np.ndarray  # column spacing (x), row spacing (y), mm
    frame_z: np.ndarray  # z of each frame, increasing, mm

    @classmethod
    def from_dataset(cls, ds: Dataset) -> DoseGrid:
        if not np.allclose(np.asarray(ds.ImageOrientationPatient, float), [1, 0, 0, 0, 1, 0]):
            raise ValueError("only axial, unrotated dose grids are handled here")
        values = ds.pixel_array.astype(np.float64) * float(ds.DoseGridScaling)
        if values.ndim == 2:
            values = values[None]
        position = np.asarray(ds.ImagePositionPatient, float)
        frame_z = position[2] + np.asarray(ds.GridFrameOffsetVector, float)
        if frame_z[0] > frame_z[-1]:
            values, frame_z = values[::-1], frame_z[::-1]
        spacing = np.asarray(ds.PixelSpacing, float)
        return cls(
            np.ascontiguousarray(values),
            position[:2],
            np.array([spacing[1], spacing[0]]),
            np.ascontiguousarray(frame_z),
        )

    def sample(self, points: np.ndarray) -> np.ndarray:
        """Dose at each ``(x, y, z)`` point; NaN outside the grid."""
        out = np.empty(len(points))
        frames = np.arange(len(self.frame_z), dtype=np.float64)
        for start in range(0, len(points), CHUNK):
            p = points[start : start + CHUNK]
            column = (p[:, 0] - self.origin_xy[0]) / self.spacing_xy[0]
            row = (p[:, 1] - self.origin_xy[1]) / self.spacing_xy[1]
            frame = np.interp(p[:, 2], self.frame_z, frames, left=-1.0, right=len(frames))
            out[start : start + len(p)] = ndimage.map_coordinates(
                self.values, [frame, row, column], order=1, mode="constant", cval=np.nan
            )
        return out


# --------------------------------------------------------------------------
# One structure against one dose, and what a method returns
# --------------------------------------------------------------------------


@dataclass
class Case:
    rtss: Dataset
    roi: int
    dose_ds: Dataset
    dose: DoseGrid
    ct: sitk.Image  # geometry only: masks read nothing else from it


@dataclass
class Result:
    metrics: dict[str, float]
    volume_at: Callable[[np.ndarray], np.ndarray]  # cumulative DVH, cc at >= dose
    seconds: float = 0.0  # one structure against one dose, from nothing
    samples: int = 0
    outside: int = 0  # samples that fell outside the dose grid
    error: str = ""
    # dicompyler only: its own histogram read with the lookup the other methods use
    corrected: dict[str, float] | None = None


def failed(error: str) -> Result:
    return Result(
        {m: math.nan for m in METRICS}, lambda d: np.full(np.shape(d), np.nan), error=error
    )


def weighted_result(dose_gy: np.ndarray, volume_cc: np.ndarray) -> Result:
    """DVH statistics of dose samples, each standing for a volume.

    ``D{x}`` is the lowest dose the hottest x of the volume receives: walk the
    samples from the hottest down and stop where their volume reaches x.
    """
    order = np.argsort(dose_gy, kind="stable")[::-1]
    dose = dose_gy[order]
    cumulative = np.cumsum(volume_cc[order])
    total = float(cumulative[-1])

    def dose_at(volume: float) -> float:
        if volume > total * (1 + 1e-12):
            return math.nan
        i = int(np.searchsorted(cumulative, volume - total * 1e-12))
        return float(dose[min(i, dose.size - 1)])

    metrics = {
        "volume_cc": total,
        "dmin": float(dose[-1]),
        "dmax": float(dose[0]),
        "dmean": float(np.dot(dose_gy, volume_cc) / total),
    }
    for x in (99, 95, 5, 1):
        metrics[f"d{x}"] = dose_at(total * x / 100)
    metrics["d0.03cc"] = dose_at(0.03)
    descending = -dose

    def volume_at(doses: np.ndarray) -> np.ndarray:
        n = np.searchsorted(descending, -np.asarray(doses, float), side="right")
        return np.where(n > 0, cumulative[np.maximum(n - 1, 0)], 0.0)

    return Result(metrics, volume_at, samples=int(dose.size))


# --------------------------------------------------------------------------
# The methods
# --------------------------------------------------------------------------


def dicompyler_result(case: Case, *, supersample: bool) -> Result:
    """dicompyler-core, as production calls it or supersampled in-plane."""
    from dicompylercore import dvhcalc

    kwargs = {
        "calculate_full_volume": True,
        "thickness": _single_plane_thickness(case.rtss, case.dose_ds, case.roi),
    }
    if supersample:
        kwargs["interpolation_resolution"] = _supersample_resolution(case.dose_ds)
    dvh = dvhcalc.get_dvh(case.rtss, case.dose_ds, case.roi, **kwargs)
    if not supersample and not dvh.volume:
        # Production retries a structure that missed every dose-grid point.
        dvh = dvhcalc.get_dvh(
            case.rtss,
            case.dose_ds,
            case.roi,
            interpolation_resolution=_supersample_resolution(case.dose_ds),
            **kwargs,
        )
    if not dvh.volume:
        raise ValueError("dicompyler-core found no volume")
    metrics = {
        "volume_cc": float(dvh.volume),
        "dmin": float(dvh.min),
        "dmax": float(dvh.max),
        "dmean": float(dvh.mean),
    }
    for x in (99, 95, 5, 1):
        metrics[f"d{x}"] = _statistic_value(dvh, f"D{x}")
    metrics["d0.03cc"] = _statistic_value(dvh, "D0.03cc")
    if not supersample:
        # This is meant to be production's number, not a re-implementation of it.
        production = compute_dvh_metrics(case.rtss, case.dose_ds, case.roi, CONFIG)
        differ = [
            name
            for key, name in PRODUCTION_KEY.items()
            if not math.isclose(production[key], metrics[name], rel_tol=1e-12, abs_tol=1e-12)
            and not (math.isnan(production[key]) and math.isnan(metrics[name]))
        ]
        if differ:
            raise AssertionError(f"differs from compute_dvh_metrics on {differ}")
    counts = np.asarray(dvh.cumulative.counts, float)
    width = float(dvh.bins[1] - dvh.bins[0])

    def volume_at(doses: np.ndarray) -> np.ndarray:
        # dicompyler's own lookup (volume_constraint): the nearest dose bin.
        index = np.rint(np.asarray(doses, float) / width).astype(np.int64)
        return np.where(index < counts.size, counts[np.clip(index, 0, counts.size - 1)], 0.0)

    # The same histogram, with D{x} read as the other methods read it, so the
    # effect of dicompyler's lookup can be told apart from that of its sampling.
    differential = dvh.differential
    bin_volume = np.asarray(differential.counts, float)
    filled = bin_volume > 0
    read = weighted_result(np.asarray(differential.bincenters, float)[filled], bin_volume[filled])
    corrected = {k: read.metrics[k] for k in ("d99", "d95", "d5", "d1", "d0.03cc")}
    return Result(metrics, volume_at, corrected=corrected)


def odd_factor(spacing_mm: float, target_mm: float | None) -> int:
    """Sub-samples per voxel edge: the fewest reaching ``target_mm``, and odd.

    Odd, so one sub-sample sits on the voxel centre, where the unsampled
    methods put their only one (PlanIQ supersamples by odd factors too).
    """
    if not target_mm:
        return 1
    k = max(1, math.ceil(spacing_mm / target_mm - 1e-9))
    return k if k % 2 else k + 1


def sub_offsets(factors: list[int]) -> np.ndarray:
    """Sub-sample positions inside one voxel, ``(m, 3)`` in voxel units."""
    axes = [(np.arange(k) + 0.5) / k - 0.5 for k in factors]
    gx, gy, gz = np.meshgrid(*axes, indexing="ij")
    return np.stack([gx.ravel(), gy.ravel(), gz.ravel()], axis=1)


def index_to_world(image: sitk.Image, index: np.ndarray) -> np.ndarray:
    origin = np.asarray(image.GetOrigin(), float)
    spacing = np.asarray(image.GetSpacing(), float)
    direction = np.asarray(image.GetDirection(), float).reshape(3, 3)
    return origin + (index * spacing) @ direction.T


class DoseHistogram:
    """A differential DVH built up sample by sample, in bins of ``BIN_GY``.

    Memory stays at one array of bins however many samples a structure takes,
    so a large structure can be sampled finely a slice at a time. Dmin, Dmax
    and Dmean are kept exactly; D{x} is read to the centre of its bin, within
    half a bin (0.5 mGy) of the sample it stands for.
    """

    def __init__(self) -> None:
        self.volume = np.zeros(1 << 16)
        self.total = 0.0
        self.dose_volume = 0.0
        self.low = math.inf
        self.high = -math.inf
        self.samples = 0
        self.outside = 0

    def add(self, dose: np.ndarray, volume_cc: np.ndarray) -> None:
        inside = np.isfinite(dose)
        self.samples += int(dose.size)
        self.outside += int(dose.size - inside.sum())
        dose, volume_cc = dose[inside], volume_cc[inside]
        if not dose.size:
            return
        index = np.maximum(np.floor(dose / BIN_GY), 0).astype(np.int64)
        top = int(index.max()) + 1
        if top > self.volume.size:
            grown = np.zeros(max(top, 2 * self.volume.size))
            grown[: self.volume.size] = self.volume
            self.volume = grown
        self.volume[:top] += np.bincount(index, weights=volume_cc, minlength=top)
        self.total += float(volume_cc.sum())
        self.dose_volume += float(np.dot(dose, volume_cc))
        self.low = min(self.low, float(dose.min()))
        self.high = max(self.high, float(dose.max()))

    def result(self) -> Result:
        if self.total <= 0:
            raise ValueError("no sample lies inside the dose grid")
        filled = np.nonzero(self.volume)[0]
        result = weighted_result((filled + 0.5) * BIN_GY, self.volume[filled])
        result.metrics.update(
            volume_cc=self.total,
            dmin=self.low,
            dmax=self.high,
            dmean=self.dose_volume / self.total,
        )
        result.samples, result.outside = self.samples, self.outside
        return result


def mask_dvh(case: Case, subsample_mm: float | None) -> Result:
    """The production mask's voxels, each split into sub-samples if asked.

    Streamed a slice at a time into a :class:`DoseHistogram`.
    """
    image, _notes = mask_with_reading(case.ct, case.rtss, case.roi)
    mask = sitk.GetArrayViewFromImage(image)
    spacing = np.asarray(case.ct.GetSpacing(), float)
    offsets = sub_offsets([odd_factor(s, subsample_mm) for s in spacing])
    weight = float(np.prod(spacing)) / 1000.0 / len(offsets)
    per_chunk = max(1, CHUNK // len(offsets))
    histogram = DoseHistogram()
    for z in np.nonzero(mask.any(axis=(1, 2)))[0]:
        y, x = np.nonzero(mask[z])
        centres = np.stack([x, y, np.full(x.size, z)], axis=1).astype(np.float64)
        for start in range(0, len(centres), per_chunk):
            points = (centres[start : start + per_chunk, None, :] + offsets[None]).reshape(-1, 3)
            dose = case.dose.sample(index_to_world(case.ct, points))
            histogram.add(dose, np.full(len(points), weight))
    return histogram.result()


def _oriented_rings(region) -> list[np.ndarray]:
    """Every ring of a region, outer boundaries anticlockwise and holes clockwise."""
    rings = []
    for part in getattr(region, "geoms", [region]):
        part = orient(part, sign=1.0)
        rings.append(np.asarray(part.exterior.coords))
        rings.extend(np.asarray(hole.coords) for hole in part.interiors)
    return rings


def edge_pieces(rings: list[np.ndarray]) -> tuple[np.ndarray, ...]:
    """The rings' edges cut at every unit grid line: ``(xa, ya, xb, yb)`` per piece.

    Each piece then lies within one cell ``[i, i+1] x [j, j+1]``.
    """
    x0 = np.concatenate([r[:-1, 0] for r in rings])
    y0 = np.concatenate([r[:-1, 1] for r in rings])
    x1 = np.concatenate([r[1:, 0] for r in rings])
    y1 = np.concatenate([r[1:, 1] for r in rings])
    n = x0.size

    def cuts(a0: np.ndarray, a1: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
        lo = np.floor(np.minimum(a0, a1)) + 1  # grid lines strictly inside the edge
        hi = np.ceil(np.maximum(a0, a1)) - 1
        count = np.maximum(hi - lo + 1, 0).astype(np.int64)
        edge = np.repeat(np.arange(n), count)
        line = lo[edge] + (np.arange(int(count.sum())) - np.repeat(np.cumsum(count) - count, count))
        return edge, (line - a0[edge]) / (a1[edge] - a0[edge])

    ex, tx = cuts(x0, x1)
    ey, ty = cuts(y0, y1)
    edge = np.concatenate([np.arange(n), np.arange(n), ex, ey])
    t = np.concatenate([np.zeros(n), np.ones(n), tx, ty])
    order = np.lexsort((t, edge))
    edge, t = edge[order], t[order]
    same = edge[1:] == edge[:-1]
    e, ta, tb = edge[1:][same], t[:-1][same], t[1:][same]
    dx, dy = x1[e] - x0[e], y1[e] - y0[e]
    return x0[e] + ta * dx, y0[e] + ta * dy, x0[e] + tb * dx, y0[e] + tb * dy


def polygon_cells(region, fx: int, fy: int) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """The sub-cells a region covers: where to sample each (voxel units), and its share.

    Sub-cells are ``1/fx`` by ``1/fy`` of a voxel, aligned to the voxel edges.
    Each is weighted by the exact area of the region inside it and sampled at
    the centroid of that area, found without clipping by the signed-area
    accumulation fonts are rasterised with. The outline is cut at every grid
    line; a piece adds the signed area between itself and its cell's right side,
    and its full height to every cell to its right in the same row, so a running
    sum along each row gives every cell its covered area. The same sums of
    first moments give the centroids. Outer boundaries run anticlockwise and
    holes clockwise, so a hole subtracts. The work grows with the outline's
    length, not the region's area.
    """
    minx, miny, maxx, maxy = region.bounds
    ix0, iy0 = math.floor((minx + 0.5) * fx), math.floor((miny + 0.5) * fy)
    nx = math.ceil((maxx + 0.5) * fx) - ix0
    ny = math.ceil((maxy + 0.5) * fy) - iy0
    scale, shift = np.array([fx, fy], float), np.array([ix0, iy0], float)
    unit = shapely.transform(region, lambda c: (c + 0.5) * scale - shift)  # cell i: [i, i+1]
    xa, ya, xb, yb = edge_pieces(_oriented_rings(unit))
    i = np.clip(np.floor(0.5 * (xa + xb)).astype(np.int64), 0, nx - 1)
    j = np.clip(np.floor(0.5 * (ya + yb)).astype(np.int64), 0, ny - 1)
    # Each piece relative to its own cell's corner, so every term is of order one
    # and a sliver's centroid does not drown in cancellation.
    xa, xb, ya, yb = xa - i, xb - i, ya - j, yb - j
    dy = yb - ya
    band_y = 0.5 * (yb * yb - ya * ya)  # first moment in y of the piece's horizontal band
    # Within the piece's own cell: the part between the piece and the cell's right side.
    own_area = dy * (1.0 - 0.5 * (xa + xb))
    own_x = 0.5 * dy * (1.0 - (xa * xa + xa * xb + xb * xb) / 3.0)
    own_y = band_y - dy * (2 * ya * xa + ya * xb + yb * xa + 2 * yb * xb) / 6.0
    grids = np.zeros((5, ny, nx))
    for k, values in enumerate((own_area, own_x, own_y, dy, band_y)):
        np.add.at(grids[k], (j, i), values)
    own_area, own_x, own_y, height, height_y = grids
    # Every cell to the right in the row gets the piece's whole band.
    height = np.cumsum(height, axis=1) - height
    height_y = np.cumsum(height_y, axis=1) - height_y
    area = -(own_area + height)
    moment_x = -(own_x + 0.5 * height)
    moment_y = -(own_y + height_y)
    covered = area > SLIVER
    share = area[covered]
    row, column = np.nonzero(covered)
    x = (column + moment_x[covered] / share + ix0) / fx - 0.5
    y = (row + moment_y[covered] / share + iy0) / fy - 0.5
    return x, y, np.minimum(share, 1.0)


def polygon_dvh(case: Case, subsample_mm: float | None) -> Result:
    """The shared reading's regions, weighted by the exact area of each sub-cell.

    Through-plane each region fills its slice (the slab every method assumes)
    and is sampled at the same sub-slab positions as ``mask-ss``. Streamed a
    slice at a time into a :class:`DoseHistogram`.
    """
    item = _find_roi_contour(case.rtss, case.roi)
    _volume, reading = _fill_structure(case.ct, item)
    spacing = np.asarray(case.ct.GetSpacing(), float)
    fx, fy, fz = (odd_factor(s, subsample_mm) for s in spacing)
    z_offsets = (np.arange(fz) + 0.5) / fz - 0.5
    cell_cc = float(np.prod(spacing)) / (fx * fy * fz) / 1000.0
    histogram = DoseHistogram()
    for z_index, region in reading.regions.items():
        if region.is_empty:
            continue
        x, y, coverage = polygon_cells(region, fx, fy)
        for start in range(0, x.size, CHUNK):
            xs, ys = x[start : start + CHUNK], y[start : start + CHUNK]
            weight = coverage[start : start + CHUNK] * cell_cc
            for dz in z_offsets:
                points = np.stack([xs, ys, np.full(xs.size, z_index + dz)], axis=1)
                histogram.add(case.dose.sample(index_to_world(case.ct, points)), weight)
    return histogram.result()


def run_method(method: str, case: Case, subsample_mm: float | None = None) -> Result:
    """One method on one structure and dose, timed from nothing; a failure is a result."""
    start = time.perf_counter()
    try:
        if method in ("dicompyler", "dicompyler-ss"):
            result = dicompyler_result(case, supersample=method == "dicompyler-ss")
        elif method == "mask":
            result = mask_dvh(case, None)
        elif method == "mask-ss":
            result = mask_dvh(case, subsample_mm)
        elif method == "polygon":
            result = polygon_dvh(case, subsample_mm)
        else:
            raise ValueError(f"unknown method {method}")
    except Exception as exc:  # noqa: BLE001 — a failure is a result to report
        result = failed(f"{type(exc).__name__}: {exc}")
    result.seconds = time.perf_counter() - start
    return result


# --------------------------------------------------------------------------
# Nelms et al. 2015
# --------------------------------------------------------------------------

NELMS_CT = {"0.2mm": "0.2 mm CT", "1mm": "1 mm CT", "2mm": "2 mm CT", "3mm": "3 mm CT"}
NELMS_DOSE = {
    "0.4x0.2x0.4": "0-4_0-2_0-4_mm",
    "1": "1mm",
    "2": "2mm",
    "3": "3mm",
}
NELMS_GRADIENT = {"AP": "Linear_AntPost", "SI": "Linear_SupInf"}
# Where each shape's tabulated analytic DVH lives in "Tabulated analytical DVHs.xlsx":
# sheet, then the dose column and first volume column for SI and for AP.
NELMS_CURVES = {
    "Sphere": ("Sphere DVH", (0, 1), (5, 6)),
    "Cylinder": ("Cylinder DVH", (0, 1), (5, 6)),
    "RtCylinder": ("Cylinder DVH", (36, 37), (41, 42)),
    "Cone": ("Cone DVH", (0, 1), (5, 6)),
    "RtCone": ("Cone DVH", (36, 37), (41, 42)),
}
NELMS_SPACING_ORDER = {"0.2mm": 0, "1mm": 1, "2mm": 2, "3mm": 3}
TEST_NAMES = {
    "1": "Test 1: contours every 0.2 mm, dose grid 0.4-3 mm",
    "2": "Test 2: contours and dose grid both 1, 2 or 3 mm, aligned",
    "2s": "Test 2, shifted half a dose voxel off the grid",
}

# The paper's results (Tables I-III), quoted beside ours for scale.
# Per parameter: (n beyond 3 %, lowest %, highest %).
PUBLISHED = {
    "1": {
        "Pinnacle3": {
            "volume_cc": (0, -2.0, 1.9),
            "dmin": (20, -7.5, 2.6),
            "dmax": (0, -1.1, 1.1),
            "dmean": (0, -1.9, 0.0),
            "d99": (20, -1.4, 7.5),
            "d95": (12, -6.6, 5.2),
            "d5": (0, -1.8, 1.0),
            "d1": (0, -0.9, 2.2),
            "d0.03cc": (0, -0.9, 1.3),
        },
        "PlanIQ": {
            "volume_cc": (0, -0.4, 0.9),
            "dmin": (0, 0.0, 0.0),
            "dmax": (0, 0.0, 0.0),
            "dmean": (0, -0.1, 0.7),
            "d99": (3, -1.8, 5.2),
            "d95": (2, -0.8, 3.9),
            "d5": (0, -0.4, 0.8),
            "d1": (0, -0.4, 0.4),
            "d0.03cc": (0, -0.4, 0.3),
        },
    },
    "2": {
        "Pinnacle3": {
            "volume_cc": (0, -2.8, 2.1),
            "dmin": (30, -7.5, 60.0),
            "dmax": (10, -5.1, 1.1),
            "dmean": (0, -1.9, 0.0),
            "d99": (18, -4.2, 44.4),
            "d95": (19, -7.8, 19.5),
            "d5": (2, -3.6, 5.3),
            "d1": (7, -8.1, 2.6),
            "d0.03cc": (7, -7.8, 0.9),
        },
        "PlanIQ": {
            "volume_cc": (1, -4.2, 0.6),
            "dmin": (0, 0.0, 0.0),
            "dmax": (0, 0.0, 0.0),
            "dmean": (0, 0.0, 0.0),
            "d99": (11, -4.2, 22.3),
            "d95": (4, -2.9, 7.0),
            "d5": (0, -1.7, 0.5),
            "d1": (1, -3.9, 0.8),
            "d0.03cc": (1, -4.6, 0.9),
        },
    },
}
# Table III "Average (N = 5)": minimum, maximum, mean and SD of the volume error (%).
PUBLISHED_TEST3 = {
    ("1mm", "SI"): {"Pinnacle3": (-4.4, 2.2, -0.3, 1.3), "PlanIQ": (-0.4, 0.4, 0.0, 0.2)},
    ("1mm", "AP"): {"Pinnacle3": (-3.5, 0.8, -0.8, 1.0), "PlanIQ": (-0.4, 0.6, 0.1, 0.2)},
    ("3mm", "SI"): {"Pinnacle3": (-11.2, 7.9, -0.7, 3.5), "PlanIQ": (-1.5, 0.6, -0.5, 0.5)},
    ("3mm", "AP"): {"Pinnacle3": (-4.1, 0.8, -1.2, 1.1), "PlanIQ": (-1.5, 0.8, -0.3, 0.6)},
}


@dataclass
class NelmsRow:
    structure: str  # file stem, e.g. RtCone_30_X15
    spacing: str  # contour spacing, e.g. "3mm"
    shift: str  # "0", or the shift off the dose grid
    voxel: str  # dose voxel, e.g. "0.4x0.2x0.4" or "3"
    gradient: str  # "AP" or "SI"
    truth: dict[str, float]

    @property
    def test(self) -> str:
        if self.spacing == "0.2mm":
            return "1"
        return "2" if self.shift == "0" else "2s"

    @property
    def shape(self) -> str:
        return self.structure.split("_")[0]


def read_nelms_rows(root: Path) -> list[NelmsRow]:
    import openpyxl

    book = openpyxl.load_workbook(
        root / "ANALYTICAL RESULTS" / "Extracted Analytical D points.xlsx",
        data_only=True,
        read_only=True,
    )
    rows = []
    for values in book.worksheets[0].iter_rows(min_row=5, values_only=True):
        if not values or not values[0]:
            continue
        analytic = values[18:27]  # volume (cc), then doses (cGy)
        truth = {"volume_cc": float(analytic[0])}
        for key, value in zip(METRICS[1:], analytic[1:]):
            truth[key] = float(value) / 100.0
        rows.append(
            NelmsRow(
                structure=str(values[0]).replace("__", "_"),  # one name is misspelt
                spacing=str(values[1]),
                shift=str(values[2]),
                voxel=str(values[3]),
                gradient="AP" if str(values[4]).startswith("Z") else "SI",
                truth=truth,
            )
        )
    book.close()
    return rows


def read_nelms_curves(root: Path) -> dict[tuple[str, str, str], tuple[np.ndarray, np.ndarray]]:
    """The tabulated analytic DVHs: (shape, gradient, spacing) -> (dose Gy, volume cc)."""
    import openpyxl

    book = openpyxl.load_workbook(
        root / "ANALYTICAL RESULTS" / "Tabulated analytical DVHs.xlsx",
        data_only=True,
        read_only=True,
    )
    sheets: dict[str, list[tuple]] = {}
    curves = {}
    for shape, (sheet, si, ap) in NELMS_CURVES.items():
        if sheet not in sheets:
            sheets[sheet] = [r for r in book[sheet].iter_rows(min_row=4, values_only=True)]
        table = sheets[sheet]
        for gradient, (dose_col, first_col) in (("SI", si), ("AP", ap)):
            dose = np.array([float(r[dose_col]) for r in table if r[dose_col] is not None])
            for spacing, offset in NELMS_SPACING_ORDER.items():
                volume = np.array([float(r[first_col + offset]) for r in table[: dose.size]])
                curves[(shape, gradient, spacing)] = (dose / 100.0, volume)
    book.close()
    return curves


def nelms_structure_path(root: Path, stem: str) -> Path:
    folder = "Spheres" if "Sphere" in stem else "Cylinders" if "Cylinder" in stem else "Cones"
    return root / "STRUCTURES" / folder / f"{stem}.dcm"


def closed_roi(ds: Dataset) -> int:
    """The structure's ROI number (each file also carries a point of interest)."""
    for item in ds.ROIContourSequence:
        kinds = {str(c.ContourGeometricType) for c in getattr(item, "ContourSequence", [])}
        if kinds == {"CLOSED_PLANAR"}:
            return int(item.ReferencedROINumber)
    raise ValueError("no closed-planar ROI")


def ct_geometry(folder: Path) -> sitk.Image:
    """A blank image with the CT series' geometry, which is all a mask reads."""
    headers = [pydicom.dcmread(str(f), stop_before_pixels=True) for f in folder.glob("*.dcm")]
    headers.sort(key=lambda h: float(h.ImagePositionPatient[2]))
    z = np.array([float(h.ImagePositionPatient[2]) for h in headers])
    steps = np.diff(z)
    if not np.allclose(steps, steps[0], atol=1e-3):
        raise ValueError("unevenly spaced CT slices")
    first = headers[0]
    iop = np.asarray(first.ImageOrientationPatient, float)
    image = sitk.Image(int(first.Columns), int(first.Rows), len(headers), sitk.sitkUInt8)
    image.SetOrigin(tuple(float(v) for v in first.ImagePositionPatient))
    image.SetSpacing((float(first.PixelSpacing[1]), float(first.PixelSpacing[0]), float(steps[0])))
    normal = np.cross(iop[:3], iop[3:])
    image.SetDirection(tuple(np.column_stack([iop[:3], iop[3:], normal]).ravel()))
    return image


def curve_error(result: Result, dose: np.ndarray, truth: np.ndarray, total: float) -> np.ndarray:
    """Volume error along the DVH, as % of the true total volume."""
    return (result.volume_at(dose) - truth) / total * 100.0


def run_nelms(
    root: Path, subsample_mm: float, variants: list[tuple[str, float | None]]
) -> list[dict]:
    rows = read_nelms_rows(root)
    curves = read_nelms_curves(root)
    test3_dose = np.round(np.arange(0.0, 30.0001, 0.1), 6)  # 301 points, as published
    doses: dict[tuple[str, str], tuple[Dataset, DoseGrid]] = {}
    cts: dict[str, sitk.Image] = {}
    groups: dict[tuple[str, str], list[NelmsRow]] = defaultdict(list)
    for row in rows:
        groups[(row.structure, row.spacing)].append(row)
    records = []
    for n, ((stem, spacing), group) in enumerate(groups.items(), 1):
        print(f"  Nelms {n}/{len(groups)}: {stem}", flush=True)
        rtss = pydicom.dcmread(str(nelms_structure_path(root, stem)))
        roi = closed_roi(rtss)
        if spacing not in cts:
            cts[spacing] = ct_geometry(root / "CT" / NELMS_CT[spacing])
        for row in group:
            key = (row.voxel, row.gradient)
            if key not in doses:
                name = f"{NELMS_GRADIENT[row.gradient]}_{NELMS_DOSE[row.voxel]}_Aligned.dcm"
                ds = pydicom.dcmread(str(root / "DOSE GRIDS" / name))
                doses[key] = (ds, DoseGrid.from_dataset(ds))
            dose_ds, grid = doses[key]
            case = Case(rtss, roi, dose_ds, grid, cts[spacing])
            for method, subsample in variants:
                if subsample is not None and row.test != "2":
                    continue  # the sub-sampling sweep runs on Test 2 only
                result = run_method(method, case, subsample or subsample_mm)
                record = {
                    "benchmark": "nelms",
                    "row": row,
                    "method": method if subsample is None else f"{method} @ {subsample:g} mm",
                    "result": result,
                }
                curve_key = (row.shape, row.gradient, row.spacing)
                if row.test == "2" and row.spacing in ("1mm", "3mm") and curve_key in curves:
                    dose, volume = curves[curve_key]
                    truth = np.interp(test3_dose, dose, volume, right=0.0)
                    error = curve_error(result, test3_dose, truth, float(volume[0]))
                    record["test3"] = (
                        float(np.min(error)),
                        float(np.max(error)),
                        float(np.mean(error)),
                        float(np.std(error, ddof=1)),
                    )
                records.append(record)
    return records


# --------------------------------------------------------------------------
# Disc phantoms
# --------------------------------------------------------------------------

D0_GY = 50.0  # dose at the phantom's centre
SLOPE_GY_PER_MM = 1.0  # so a dose error in Gy is also a boundary shift in mm
DOSE_SCALING = 1e-6
DIRECTIONS = {"x": (1.0, 0.0), "z": (0.0, 1.0), "xz": (math.sqrt(0.5), math.sqrt(0.5))}
DIRECTION_LABEL = {"x": "in-plane (x)", "z": "through-plane (z)", "xz": "oblique (45° x-z)"}
# CT pixel, CT slice, dose voxel (mm)
GRIDS = {
    "CT 0.98 × 0.98 × 2 mm, dose 2 mm": (0.977, 2.0, 2.0),
    "CT 0.98 × 0.98 × 3 mm, dose 2.5 mm": (0.977, 3.0, 2.5),
}
SHAPES = ("sphere", "cylinder", "ring")
RADII = (2.5, 5.0, 10.0, 20.0)
VERTEX_SPACING_MM = 0.05
# A loop's ContourData must stay under the 64 KB an explicit-VR element can hold:
# at ~11 characters a coordinate, 1,500 vertices is about 50 KB.
MAX_VERTICES = 1500


@dataclass(frozen=True)
class Disc:
    z: float  # plane height relative to the phantom centre, mm
    r_out: float
    r_in: float = 0.0  # a hole, drawn as a second loop inside the first


def disc_stack(shape: str, radius: float, slice_mm: float, on_slice: bool) -> list[Disc]:
    """The contours of a sphere, cylinder (height 2R) or ring (hole R/2), one per slice.

    ``on_slice`` puts the centre on a slice; otherwise midway between two.
    Every plane where the solid exists is drawn, as a contouring system would.
    """
    count = int(math.ceil(radius / slice_mm)) + 1
    if on_slice:
        lattice = [k * slice_mm for k in range(-count, count + 1)]
    else:
        lattice = [(k + 0.5) * slice_mm for k in range(-count - 1, count + 1)]
    discs = []
    for z in lattice:
        if abs(z) >= radius:
            continue
        if shape == "sphere":
            r = math.sqrt(radius * radius - z * z)
            if r >= 0.5:
                discs.append(Disc(z, r))
        else:
            discs.append(Disc(z, radius, radius / 2 if shape == "ring" else 0.0))
    return discs


def circle_loop(radius: float) -> np.ndarray:
    """A polygon with the circle's exact area, vertices every 0.05 mm (0.084 at R = 20)."""
    n = min(MAX_VERTICES, max(64, math.ceil(2 * math.pi * radius / VERTEX_SPACING_MM)))
    scale = math.sqrt(2 * math.pi / (n * math.sin(2 * math.pi / n)))
    theta = 2 * math.pi * np.arange(n) / n
    return radius * scale * np.stack([np.cos(theta), np.sin(theta)], axis=1)


def _segment_area(u: np.ndarray, r: float) -> np.ndarray:
    """Area of a disc of radius ``r`` lying beyond the chord at offset ``u``."""
    u = np.clip(u, -r, r)
    return r * r * np.arccos(u / r) - u * np.sqrt(r * r - u * u)


def _segment_integral(u: np.ndarray, r: float) -> np.ndarray:
    """An antiderivative of :func:`_segment_area` in ``u``, valid on the whole line."""
    uc = np.clip(u, -r, r)
    inside = uc * _segment_area(uc, r) - (2.0 / 3.0) * (r * r - uc * uc) ** 1.5
    return np.where(u <= -r, math.pi * r * r * u, np.where(u >= r, 0.0, inside))


class DiscTruth:
    """The exact DVH of a disc stack in a linear dose.

    Each disc fills a slab one slice thick. Along x, the volume above a dose is
    the slab's thickness times a circular segment; along z, the disc's area
    times the part of the slab above a height; obliquely, the segment's area
    integrated through the slab, which has the closed form
    :func:`_segment_integral`.
    """

    def __init__(
        self,
        discs: list[Disc],
        slice_mm: float,
        direction: str,
        d0: float = D0_GY,
        slope: float = SLOPE_GY_PER_MM,
    ):
        self.discs = discs
        self.half = slice_mm / 2.0
        self.cx, self.cz = DIRECTIONS[direction]
        self.d0, self.slope = d0, slope

    def _slab(self, delta: np.ndarray, r: float, a: float, b: float) -> np.ndarray:
        if self.cz == 0:
            return (b - a) * _segment_area(delta / self.cx, r)
        if self.cx == 0:
            return math.pi * r * r * np.clip(b - np.maximum(a, delta), 0.0, b - a)
        high = _segment_integral((delta - self.cz * a) / self.cx, r)
        low = _segment_integral((delta - self.cz * b) / self.cx, r)
        return (self.cx / self.cz) * (high - low)

    def volume_at(self, dose) -> np.ndarray:
        """Volume (cc) receiving at least ``dose``."""
        delta = (np.asarray(dose, float) - self.d0) / self.slope
        total = np.zeros_like(delta)
        for disc in self.discs:
            a, b = disc.z - self.half, disc.z + self.half
            total = total + self._slab(delta, disc.r_out, a, b)
            if disc.r_in:
                total = total - self._slab(delta, disc.r_in, a, b)
        return total / 1000.0

    def metrics(self) -> dict[str, float]:
        slabs = [(d.z, 2 * self.half * math.pi * (d.r_out**2 - d.r_in**2)) for d in self.discs]
        total = sum(v for _, v in slabs) / 1000.0
        z_mean = sum(z * v for z, v in slabs) / (total * 1000.0)
        top = max(self.cx * d.r_out + self.cz * (d.z + self.half) for d in self.discs)
        bottom = min(-self.cx * d.r_out + self.cz * (d.z - self.half) for d in self.discs)
        dmin, dmax = self.d0 + self.slope * bottom, self.d0 + self.slope * top

        def dose_at(volume: float) -> float:
            if volume > total:
                return math.nan
            return optimize.brentq(
                lambda d: float(self.volume_at(d)) - volume, dmin, dmax, xtol=1e-10
            )

        out = {
            "volume_cc": total,
            "dmin": dmin,
            "dmax": dmax,
            "dmean": self.d0 + self.slope * self.cz * z_mean,
        }
        for x in (99, 95, 5, 1):
            out[f"d{x}"] = dose_at(total * x / 100)
        out["d0.03cc"] = dose_at(0.03)
        return out


def _meta(sop_class: str, sop_uid: str) -> Dataset:
    meta = Dataset()
    meta.MediaStorageSOPClassUID = sop_class
    meta.MediaStorageSOPInstanceUID = sop_uid
    meta.TransferSyntaxUID = ExplicitVRLittleEndian
    return meta


def _ds(values) -> list[str]:
    """Decimal strings for DICOM DS, well inside its 16 characters."""
    return [f"{float(v):.6f}" for v in values]


def write_disc_rtss(path: Path, discs: list[Disc], for_uid: str) -> Dataset:
    sop_uid = generate_uid()
    ds = FileDataset(
        str(path), {}, file_meta=_meta(RTStructureSetStorage, sop_uid), preamble=b"\0" * 128
    )
    ds.PatientID = "PHANTOM"
    ds.PatientName = "Disc^Phantom"
    ds.StudyInstanceUID = generate_uid()
    ds.SeriesInstanceUID = generate_uid()
    ds.SOPInstanceUID = sop_uid
    ds.SOPClassUID = RTStructureSetStorage
    ds.Modality = "RTSTRUCT"
    ds.StructureSetLabel = "Discs"
    reference = Dataset()
    reference.FrameOfReferenceUID = for_uid
    ds.ReferencedFrameOfReferenceSequence = [reference]
    roi = Dataset()
    roi.ROINumber = 1
    roi.ROIName = "discs"
    roi.ReferencedFrameOfReferenceUID = for_uid
    roi.ROIGenerationAlgorithm = ""
    ds.StructureSetROISequence = [roi]
    contours = Dataset()
    contours.ReferencedROINumber = 1
    contours.ROIDisplayColor = [255, 0, 0]
    contours.ContourSequence = []
    for disc in discs:
        for r in (disc.r_out, disc.r_in):
            if not r:
                continue
            loop = circle_loop(r)
            item = Dataset()
            item.ContourGeometricType = "CLOSED_PLANAR"
            item.NumberOfContourPoints = len(loop)
            xyz = np.column_stack([loop, np.full(len(loop), disc.z)])
            item.ContourData = _ds(xyz.ravel())
            contours.ContourSequence.append(item)
    ds.ROIContourSequence = [contours]
    observation = Dataset()
    observation.ObservationNumber = 1
    observation.ReferencedROINumber = 1
    observation.RTROIInterpretedType = "ORGAN"
    observation.ROIInterpreter = ""
    ds.RTROIObservationsSequence = [observation]
    ds.is_little_endian = True
    ds.is_implicit_VR = False
    pydicom.dcmwrite(str(path), ds, write_like_original=False)
    return pydicom.dcmread(str(path))


def write_linear_rtdose(
    path: Path,
    *,
    direction: str,
    origin: np.ndarray,
    spacing: float,
    shape: tuple,
    for_uid: str,
    d0: float = D0_GY,
    slope: float = SLOPE_GY_PER_MM,
) -> Dataset:
    """A dose of ``d0`` at the origin rising ``slope`` Gy/mm along ``direction``."""
    frames, rows, columns = shape
    x = origin[0] + spacing * np.arange(columns)
    z = origin[2] + spacing * np.arange(frames)
    cx, cz = DIRECTIONS[direction]
    dose = d0 + slope * (cx * x[None, None, :] + cz * z[:, None, None])
    dose = np.broadcast_to(dose, (frames, rows, columns))
    if dose.min() < 0:
        raise ValueError("the dose grid reaches below zero")
    sop_uid = generate_uid()
    ds = FileDataset(str(path), {}, file_meta=_meta(RTDoseStorage, sop_uid), preamble=b"\0" * 128)
    ds.PatientID = "PHANTOM"
    ds.PatientName = "Disc^Phantom"
    ds.StudyInstanceUID = generate_uid()
    ds.SeriesInstanceUID = generate_uid()
    ds.SOPInstanceUID = sop_uid
    ds.SOPClassUID = RTDoseStorage
    ds.Modality = "RTDOSE"
    ds.FrameOfReferenceUID = for_uid
    ds.ImagePositionPatient = _ds(origin)
    ds.ImageOrientationPatient = [1, 0, 0, 0, 1, 0]
    ds.PixelSpacing = _ds([spacing, spacing])
    ds.Rows = rows
    ds.Columns = columns
    ds.NumberOfFrames = frames
    ds.FrameIncrementPointer = (0x3004, 0x000C)
    ds.GridFrameOffsetVector = _ds(spacing * np.arange(frames))
    ds.SamplesPerPixel = 1
    ds.PhotometricInterpretation = "MONOCHROME2"
    ds.BitsAllocated = 32
    ds.BitsStored = 32
    ds.HighBit = 31
    ds.PixelRepresentation = 0
    ds.DoseUnits = "GY"
    ds.DoseType = "PHYSICAL"
    ds.DoseSummationType = "PLAN"
    ds.DoseGridScaling = "1e-06"
    ds.PixelData = np.round(dose / DOSE_SCALING).astype("<u4").tobytes()
    ds.is_little_endian = True
    ds.is_implicit_VR = False
    pydicom.dcmwrite(str(path), ds, write_like_original=False)
    return pydicom.dcmread(str(path))


def blank_ct(pixel_mm: float, slice_mm: float, discs: list[Disc], radius: float, rng) -> sitk.Image:
    """CT geometry with the phantom's centre at a random sub-voxel position in-plane."""
    n_xy = int(math.ceil(2 * (radius + 4) / pixel_mm)) + 2
    ux, uy = rng.random(2)
    heights = sorted(d.z for d in discs)
    n_z = int(round((heights[-1] - heights[0]) / slice_mm)) + 5
    image = sitk.Image(n_xy, n_xy, n_z, sitk.sitkUInt8)
    image.SetOrigin(
        (-(n_xy // 2 + ux) * pixel_mm, -(n_xy // 2 + uy) * pixel_mm, heights[0] - 2 * slice_mm)
    )
    image.SetSpacing((pixel_mm, pixel_mm, slice_mm))
    return image


def dose_box(
    discs: list[Disc], radius: float, slice_mm: float, dose_mm: float, offset: np.ndarray
) -> tuple[np.ndarray, tuple[int, int, int]]:
    """Origin and shape of a dose grid three voxels clear of the phantom, offset randomly."""
    reach_xy = radius + 3 * dose_mm
    reach_z = max(abs(d.z) for d in discs) + slice_mm / 2 + 3 * dose_mm
    m_xy, m_z = math.ceil(reach_xy / dose_mm), math.ceil(reach_z / dose_mm)
    origin = -np.array([m_xy + offset[0], m_xy + offset[1], m_z + offset[2]]) * dose_mm
    return origin, (2 * m_z + 2, 2 * m_xy + 2, 2 * m_xy + 2)


@dataclass
class DiscCase:
    grid: str
    shape: str
    radius: float
    direction: str
    placement: int
    case: Case
    truth: DiscTruth
    truth_metrics: dict[str, float] = field(default_factory=dict)


def disc_cases(workdir: Path, placements: int, seed: int) -> Iterator[DiscCase]:
    rng = np.random.default_rng(seed)
    for grid, (pixel_mm, slice_mm, dose_mm) in GRIDS.items():
        for shape in SHAPES:
            for radius in RADII:
                for placement in range(placements):
                    discs = disc_stack(shape, radius, slice_mm, on_slice=bool(rng.integers(2)))
                    for_uid = generate_uid()
                    ct = blank_ct(pixel_mm, slice_mm, discs, radius, rng)
                    rtss = write_disc_rtss(workdir / "rtss.dcm", discs, for_uid)
                    origin, shape_zyx = dose_box(discs, radius, slice_mm, dose_mm, rng.random(3))
                    for direction in DIRECTIONS:
                        dose_ds = write_linear_rtdose(
                            workdir / f"dose_{direction}.dcm",
                            direction=direction,
                            origin=origin,
                            spacing=dose_mm,
                            shape=shape_zyx,
                            for_uid=for_uid,
                        )
                        truth = DiscTruth(discs, slice_mm, direction)
                        case = Case(rtss, 1, dose_ds, DoseGrid.from_dataset(dose_ds), ct)
                        yield DiscCase(
                            grid, shape, radius, direction, placement, case, truth, truth.metrics()
                        )


def run_discs(subsample_mm: float, placements: int, seed: int) -> list[dict]:
    records = []
    total = len(GRIDS) * len(SHAPES) * len(RADII) * placements * len(DIRECTIONS)
    with tempfile.TemporaryDirectory() as tmp:
        for n, item in enumerate(disc_cases(Path(tmp), placements, seed), 1):
            if n % 36 == 1:
                print(f"  discs {n}/{total}", flush=True)
            dose = np.linspace(item.truth_metrics["dmin"], item.truth_metrics["dmax"], 301)
            truth_curve = item.truth.volume_at(dose)
            for method in METHODS:
                result = run_method(method, item.case, subsample_mm)
                error = curve_error(result, dose, truth_curve, item.truth_metrics["volume_cc"])
                records.append(
                    {
                        "benchmark": "discs",
                        "grid": item.grid,
                        "shape": item.shape,
                        "radius": item.radius,
                        "direction": item.direction,
                        "placement": item.placement,
                        "method": method,
                        "result": result,
                        "truth": item.truth_metrics,
                        "curve_max_pct": float(np.nanmax(np.abs(error))),
                    }
                )
    return records


# --------------------------------------------------------------------------
# Large structures: time against accuracy
# --------------------------------------------------------------------------

# From a lens-sized sphere to a body-sized cylinder, on the tender cohort's grid.
LARGE = (
    ("sphere", 10.0),
    ("sphere", 20.0),
    ("sphere", 40.0),
    ("sphere", 60.0),
    ("cylinder", 80.0),
    ("cylinder", 100.0),
)
LARGE_GRID = (0.977, 2.0, 2.0)  # CT pixel, CT slice, dose voxel (mm)
LARGE_SPACINGS = (1.0, 0.5, 0.25)
LARGE_D0, LARGE_SLOPE = 100.0, 0.25  # gentler, so the dose stays positive across 20 cm
LARGE_RETIME_S = 5.0  # a run faster than this is timed twice and the faster kept


def large_variants() -> list[tuple[str, float | None]]:
    return [("dicompyler", None), ("mask", None)] + [
        (m, s) for m in SAMPLED for s in LARGE_SPACINGS
    ]


def variant_name(method: str, spacing: float | None) -> str:
    return method if spacing is None else f"{method} @ {spacing:g} mm"


def run_large(seed: int) -> list[dict]:
    rng = np.random.default_rng(seed + 2)
    pixel_mm, slice_mm, dose_mm = LARGE_GRID
    records = []
    with tempfile.TemporaryDirectory() as tmp:
        for shape, radius in LARGE:
            discs = disc_stack(shape, radius, slice_mm, on_slice=True)
            for_uid = generate_uid()
            ct = blank_ct(pixel_mm, slice_mm, discs, radius, rng)
            rtss = write_disc_rtss(Path(tmp) / "rtss.dcm", discs, for_uid)
            origin, shape_zyx = dose_box(discs, radius, slice_mm, dose_mm, rng.random(3))
            dose_ds = write_linear_rtdose(
                Path(tmp) / "dose.dcm",
                direction="xz",
                origin=origin,
                spacing=dose_mm,
                shape=shape_zyx,
                for_uid=for_uid,
                d0=LARGE_D0,
                slope=LARGE_SLOPE,
            )
            case = Case(rtss, 1, dose_ds, DoseGrid.from_dataset(dose_ds), ct)
            truth = DiscTruth(discs, slice_mm, "xz", LARGE_D0, LARGE_SLOPE).metrics()
            print(f"  large: {shape} R {radius:g} mm, {truth['volume_cc']:.0f} cc", flush=True)
            for method, spacing in large_variants():
                result = run_method(method, case, spacing)
                if result.seconds < LARGE_RETIME_S:
                    result = min(result, run_method(method, case, spacing), key=lambda r: r.seconds)
                records.append(
                    {
                        "benchmark": "large",
                        "shape": shape,
                        "radius": radius,
                        "method": variant_name(method, spacing),
                        "result": result,
                        "truth": truth,
                    }
                )
    return records


# --------------------------------------------------------------------------
# Checks on the harness itself
# --------------------------------------------------------------------------


def check_truth(seed: int) -> str:
    """The closed form against brute-force integration over the written polygons."""
    rng = np.random.default_rng(seed + 1)
    nodes, weights = np.polynomial.legendre.leggauss(96)
    worst = 0.0
    with tempfile.TemporaryDirectory() as tmp:
        for shape in ("sphere", "ring"):
            for direction in DIRECTIONS:
                discs = disc_stack(shape, 7.0, 3.0, on_slice=bool(rng.integers(2)))
                rtss = write_disc_rtss(Path(tmp) / "rtss.dcm", discs, generate_uid())
                regions: dict[float, shapely.Geometry] = {}
                for item in rtss.ROIContourSequence[0].ContourSequence:
                    xyz = np.asarray(item.ContourData, float).reshape(-1, 3)
                    z = round(float(xyz[0, 2]), 6)
                    loop = shapely.Polygon(xyz[:, :2])
                    regions[z] = regions[z].symmetric_difference(loop) if z in regions else loop
                truth = DiscTruth(discs, 3.0, direction)
                metrics = truth.metrics()
                cx, cz = DIRECTIONS[direction]
                for dose in np.linspace(metrics["dmin"] + 0.3, metrics["dmax"] - 0.3, 7):
                    delta = (dose - D0_GY) / SLOPE_GY_PER_MM
                    brute = 0.0
                    for z, region in regions.items():
                        if cx == 0:
                            # A step through the slab: no quadrature, just its height.
                            brute += region.area * float(
                                np.clip(z + 1.5 - max(z - 1.5, delta), 0, 3)
                            )
                            continue
                        for h, w in zip(z + 1.5 * nodes, weights):
                            cut = shapely.box((delta - cz * h) / cx, -1e3, 1e3, 1e3)
                            brute += w * 1.5 * region.intersection(cut).area
                    exact = float(truth.volume_at(dose)) * 1000.0
                    worst = max(worst, abs(exact - brute) / (metrics["volume_cc"] * 1000.0))
    return f"{worst:.1e}"


def _clipped_cells(region, f: int) -> dict[tuple[int, int], tuple[float, float, float]]:
    """Every sub-cell clipped against the region with shapely: the slow, obvious way."""
    h = 1.0 / f
    minx, miny, maxx, maxy = region.bounds
    x0 = np.arange(math.floor((minx + 0.5) * f), math.ceil((maxx + 0.5) * f)) / f - 0.5
    y0 = np.arange(math.floor((miny + 0.5) * f), math.ceil((maxy + 0.5) * f)) / f - 0.5
    gx, gy = np.meshgrid(x0, y0)
    gx, gy = gx.ravel(), gy.ravel()
    pieces = shapely.intersection(shapely.box(gx, gy, gx + h, gy + h), region)
    share = shapely.area(pieces) / (h * h)
    kept = share > SLIVER
    centroid = shapely.get_coordinates(shapely.centroid(pieces[kept]))
    cells = zip(gx[kept], gy[kept], centroid[:, 0], centroid[:, 1], share[kept])
    return {(round((x + 0.5) * f), round((y + 0.5) * f)): (cx, cy, s) for x, y, cx, cy, s in cells}


def check_polygon_cells(root: Path | None, seed: int) -> str:
    """Polygon's coverage by accumulation against clipping every sub-cell with shapely."""
    readings = []
    if root is not None:
        for stem, spacing in (
            ("Sphere_10_0", "1mm"),
            ("RtCylinder_30_X15Z15", "3mm"),
            ("RtCone_20_X10", "2mm"),
            ("Cone_02_0", "0.2mm"),
        ):
            rtss = pydicom.dcmread(str(nelms_structure_path(root, stem)))
            ct = ct_geometry(root / "CT" / NELMS_CT[spacing])
            readings.append(_fill_structure(ct, _find_roi_contour(rtss, closed_roi(rtss)))[1])
    rng = np.random.default_rng(seed + 3)
    with tempfile.TemporaryDirectory() as tmp:
        for shape, radius in (("ring", 2.5), ("ring", 10.0), ("sphere", 5.0)):
            discs = disc_stack(shape, radius, 3.0, on_slice=bool(rng.integers(2)))
            ct = blank_ct(0.977, 3.0, discs, radius, rng)
            rtss = write_disc_rtss(Path(tmp) / "rtss.dcm", discs, generate_uid())
            readings.append(_fill_structure(ct, _find_roi_contour(rtss, 1))[1])
    cells = mismatched = 0
    worst_share = worst_centroid = 0.0
    for reading in readings:
        for region in reading.regions.values():
            for f in (1, 3, 5):
                clipped = _clipped_cells(region, f)
                x, y, share = polygon_cells(region, f, f)
                ours = {
                    (math.floor((a + 0.5) * f), math.floor((b + 0.5) * f)): (a, b, s)
                    for a, b, s in zip(x, y, share)
                }
                cells += len(clipped)
                if ours.keys() != clipped.keys():
                    mismatched += 1
                    continue
                for key, (cx, cy, s) in clipped.items():
                    a, b, t = ours[key]
                    worst_share = max(worst_share, abs(s - t))
                    worst_centroid = max(worst_centroid, abs(cx - a), abs(cy - b))
    if mismatched:
        return f"{mismatched} slices covered different sub-cells (of {cells:,} sub-cells)"
    return (
        f"{cells:,} sub-cells, the same cells covered; largest difference in the share "
        f"covered {worst_share:.0e}, in the centroid {worst_centroid:.0e} voxel"
    )


def check_sampler(root: Path | None) -> str:
    """Our trilinear sampling against SimpleITK's, which production uses today."""
    if root is None:
        return "not run (no Nelms data)"
    rtss = pydicom.dcmread(str(nelms_structure_path(root, "Sphere_10_0")))
    ct = ct_geometry(root / "CT" / NELMS_CT["1mm"])
    dose_ds = pydicom.dcmread(str(root / "DOSE GRIDS" / "Linear_SupInf_2mm_Aligned.dcm"))
    grid = DoseGrid.from_dataset(dose_ds)
    mask, _notes = mask_with_reading(ct, rtss, closed_roi(rtss))
    image = sitk.GetImageFromArray(grid.values)
    image.SetOrigin((*grid.origin_xy, grid.frame_z[0]))
    image.SetSpacing((*grid.spacing_xy, float(grid.frame_z[1] - grid.frame_z[0])))
    resampled = sitk.GetArrayFromImage(
        sitk.Resample(image, mask, sitk.Transform(), sitk.sitkLinear, 0.0, sitk.sitkFloat64)
    )
    inside = sitk.GetArrayViewFromImage(mask) > 0
    z, y, x = np.nonzero(inside)
    ours = grid.sample(index_to_world(mask, np.stack([x, y, z], axis=1).astype(float)))
    return f"{float(np.max(np.abs(ours - resampled[inside]))):.1e} Gy"


# --------------------------------------------------------------------------
# Report
# --------------------------------------------------------------------------


def pct(value: float, truth: float) -> float:
    return (value - truth) / truth * 100.0


def table(headers: list[str], rows: list[list[str]], align: str | None = None) -> list[str]:
    align = align or "l" + "r" * (len(headers) - 1)
    rule = ["---:" if a == "r" else "---" for a in align]
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(rule) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in rows]
    return lines + [""]


def nelms_counts(records: list[dict], test: str, method: str, threshold: float = 3.0) -> dict:
    """Per parameter: (n beyond threshold, N, lowest %, highest %), as the paper counts.

    Volume is counted once per structure and dose grid: it does not depend on
    the gradient's direction.
    """
    out = {}
    chosen = [r for r in records if r["row"].test == test and r["method"] == method]
    for metric in METRICS:
        seen = set()
        deltas = []
        for r in chosen:
            if metric == "volume_cc":
                key = (r["row"].structure, r["row"].voxel)
                if key in seen:
                    continue
                seen.add(key)
            deltas.append(pct(r["result"].metrics[metric], r["row"].truth[metric]))
        d = np.array(deltas)
        finite = d[np.isfinite(d)]
        out[metric] = (
            int(np.sum(np.abs(finite) > threshold)) + int(np.sum(~np.isfinite(d))),
            len(d),
            float(finite.min()) if finite.size else math.nan,
            float(finite.max()) if finite.size else math.nan,
        )
    return out


def _range(lo: float, hi: float) -> str:
    return f"{lo:+.1f} to {hi:+.1f}"


def nelms_section(records: list[dict], test: str) -> list[str]:
    counts = {m: nelms_counts(records, test, m) for m in METHODS}
    published = PUBLISHED.get(test, {})
    headers = ["Parameter", *METHODS, *(f"{s} (paper)" for s in published)]
    rows = []
    for metric in METRICS:
        row = [LABEL[metric]]
        for m in METHODS:
            n, total, lo, hi = counts[m][metric]
            row.append(f"{n}/{total} ({_range(lo, hi)})")
        for system in published.values():
            n, lo, hi = system[metric]
            row.append(f"{n} ({_range(lo, hi)})")
        rows.append(row)
    for label, subset in (
        ("**All**", METRICS),
        ("**Without Dmin, Dmax**", ("volume_cc", *CLINICAL)),
    ):
        row = [label]
        for m in METHODS:
            n = sum(counts[m][k][0] for k in subset)
            total = sum(counts[m][k][1] for k in subset)
            row.append(f"**{n}/{total}**")
        for system in published.values():
            row.append(f"**{sum(system[k][0] for k in subset)}**")
        rows.append(row)
    lines = [f"### {TEST_NAMES[test]}", ""]
    lines += [
        "Cells: parameters more than 3 % from the analytic value / parameters scored "
        "(lowest to highest % difference). Differences are relative to the local "
        "analytic value, as published, so the low-dose parameters (Dmin, D99, D95) "
        "are amplified."
        + (" The paper columns are its Table I." if test == "1" else "")
        + (" The paper columns are its Table II." if test == "2" else ""),
        "",
    ]
    return lines + table(headers, rows)


def nelms_two_percent(records: list[dict]) -> list[str]:
    rows = []
    for test in ("1", "2", "2s"):
        row = [TEST_NAMES[test].split(":")[0]]
        for m in METHODS:
            c = nelms_counts(records, test, m, threshold=2.0)
            n = sum(c[k][0] for k in ("volume_cc", *CLINICAL))
            total = sum(c[k][1] for k in ("volume_cc", *CLINICAL))
            row.append(f"{n}/{total}")
        rows.append(row)
    return table(["Parameters beyond 2 %, without Dmin, Dmax", *METHODS], rows)


def nelms_test3(records: list[dict]) -> list[str]:
    lines = ["### Test 3: volume error along the whole curve", ""]
    lines += [
        "The volume error at every 0.1 Gy from 0 to 30 Gy (301 points), as % of the "
        "structure's analytic volume, on the Test 2 data at 1 and 3 mm. Each cell "
        "averages the five shapes' lowest, highest and mean error and SD, like the "
        "paper's Table III: lowest to highest (mean ± SD).",
        "",
    ]
    rows = []
    for spacing in ("1mm", "3mm"):
        for gradient, label in (("SI", "superior-inferior"), ("AP", "anterior-posterior")):
            row = [f"{spacing[0]} mm, {label}"]
            for m in METHODS:
                stats = np.array(
                    [
                        r["test3"]
                        for r in records
                        if r["method"] == m
                        and "test3" in r
                        and r["row"].spacing == spacing
                        and r["row"].gradient == gradient
                    ]
                )
                lo, hi, mean, sd = stats.mean(axis=0)
                row.append(f"{lo:+.1f} to {hi:+.1f} ({mean:+.1f} ± {sd:.1f})")
            for system in ("Pinnacle3", "PlanIQ"):
                lo, hi, mean, sd = PUBLISHED_TEST3[(spacing, gradient)][system]
                row.append(f"{lo:+.1f} to {hi:+.1f} ({mean:+.1f} ± {sd:.1f})")
            rows.append(row)
    return lines + table(["Data", *METHODS, "Pinnacle3 (paper)", "PlanIQ (paper)"], rows)


def nelms_test3_detail(records: list[dict]) -> list[str]:
    rows = []
    keys = sorted(
        {(r["row"].spacing, r["row"].gradient, r["row"].structure) for r in records if "test3" in r}
    )
    for _spacing, gradient, structure in keys:
        row = [f"{structure}, {gradient}"]
        for m in METHODS:
            (stats,) = [
                r["test3"]
                for r in records
                if r["method"] == m
                and "test3" in r
                and r["row"].structure == structure
                and r["row"].gradient == gradient
            ]
            row.append(f"{stats[0]:+.1f} to {stats[1]:+.1f}")
        rows.append(row)
    return table(["Dataset", *METHODS], rows)


def sweep_section(records: list[dict], main_mm: float) -> list[str]:
    lines = ["## Sub-sample spacing: accuracy against cost", ""]
    lines += [
        "The two sub-sampled methods at three spacings, on the 30 Test 2 datasets "
        "(Dmin and Dmax set aside). The sample count is per structure, and so is the "
        "time: building the samples once plus one dose lookup.",
        "",
    ]
    rows = []
    for method in SAMPLED:
        for spacing in sorted({1.0, 0.5, main_mm}, reverse=True):
            name = method if spacing == main_mm else f"{method} @ {spacing:g} mm"
            chosen = [r for r in records if r["method"] == name and r["row"].test == "2"]
            if not chosen:
                continue
            c = nelms_counts(records, "2", name)
            c2 = nelms_counts(records, "2", name, threshold=2.0)
            subset = ("volume_cc", *CLINICAL)
            n3 = sum(c[k][0] for k in subset)
            n2 = sum(c2[k][0] for k in subset)
            total = sum(c[k][1] for k in subset)
            worst = max(max(abs(c[k][2]), abs(c[k][3])) for k in subset)
            samples = np.median([r["result"].samples for r in chosen])
            seconds = np.median([r["result"].seconds for r in chosen])
            rows.append(
                [
                    f"{method} @ {spacing:g} mm",
                    f"{n3}/{total}",
                    f"{n2}/{total}",
                    f"{worst:.1f}",
                    f"{samples:,.0f}",
                    f"{seconds:.2f}",
                ]
            )
    return lines + table(["Method", "> 3 %", "> 2 %", "Worst |%|", "Samples", "Seconds"], rows)


def disc_errors(records: list[dict], method: str, **where) -> list[dict]:
    out = []
    for r in records:
        if r["method"] != method or any(r[k] != v for k, v in where.items()):
            continue
        m, t = r["result"].metrics, r["truth"]
        errors = {k: m[k] - t[k] for k in METRICS if k != "volume_cc"}
        errors["volume_pct"] = pct(m["volume_cc"], t["volume_cc"])
        errors["clinical"] = max(abs(errors[k]) for k in CLINICAL)
        errors["curve"] = r["curve_max_pct"]
        out.append(errors)
    return out


def _mean_max(values: list[float], digits: int = 2) -> str:
    v = np.abs(np.array(values, float))
    if not v.size or np.isnan(v).all():
        return "n/a"
    bad = int(np.isnan(v).sum())
    text = f"{np.nanmean(v):.{digits}f} / {np.nanmax(v):.{digits}f}"
    return text + (f" ({bad} failed)" if bad else "")


def _median_max(values: list[float], digits: int = 2) -> str:
    v = np.abs(np.array(values, float))
    if not v.size or np.isnan(v).all():
        return "n/a"
    bad = int(np.isnan(v).sum())
    text = f"{np.nanmedian(v):.{digits}f} / {np.nanmax(v):.{digits}f}"
    return text + (f" ({bad} failed)" if bad else "")


def discs_section(records: list[dict], placements: int) -> list[str]:
    n_cases = len(records) // len(METHODS)
    lines = ["## Disc phantoms", ""]
    lines += [
        f"{n_cases} cases: {len(SHAPES)} shapes × {len(RADII)} radii × {len(GRIDS)} grids × "
        f"{placements} random placements × {len(DIRECTIONS)} dose directions. The dose is "
        f"{D0_GY:g} Gy at the centre with a {SLOPE_GY_PER_MM:g} Gy/mm gradient, so an error "
        "in Gy is also the boundary shift, in mm, that would cause it.",
        "",
        "### Every case",
        "",
        "Mean / largest absolute error over all cases. Doses in Gy; volume in %.",
        "",
    ]
    errors = {m: disc_errors(records, m) for m in METHODS}
    rows = []
    for metric in METRICS:
        key = "volume_pct" if metric == "volume_cc" else metric
        rows.append(
            [LABEL[metric] + (" (%)" if metric == "volume_cc" else "")]
            + [_mean_max([e[key] for e in errors[m]]) for m in METHODS]
        )
    rows.append(
        ["Worst of Dmean-D0.03cc"]
        + [_mean_max([e["clinical"] for e in errors[m]]) for m in METHODS]
    )
    rows.append(
        ["Whole curve, largest ΔV (%)"]
        + [_mean_max([e["curve"] for e in errors[m]]) for m in METHODS]
    )
    lines += table(["Metric", *METHODS], rows)

    lines += [
        "### By size",
        "",
        "Median / largest of each case's worst error among Dmean, D99, D95, D5, D1 and "
        "D0.03cc (Gy), then the volume error (%).",
        "",
    ]
    rows = []
    for radius in RADII:
        rows.append(
            [f"R = {radius:g} mm, worst dose"]
            + [
                _median_max([e["clinical"] for e in disc_errors(records, m, radius=radius)])
                for m in METHODS
            ]
        )
    for radius in RADII:
        rows.append(
            [f"R = {radius:g} mm, volume (%)"]
            + [
                _median_max([e["volume_pct"] for e in disc_errors(records, m, radius=radius)])
                for m in METHODS
            ]
        )
    lines += table(["Size", *METHODS], rows)

    lines += ["### By dose direction and grid", "", "Median / largest worst dose error (Gy).", ""]
    rows = []
    for direction in DIRECTIONS:
        rows.append(
            [DIRECTION_LABEL[direction]]
            + [
                _median_max([e["clinical"] for e in disc_errors(records, m, direction=direction)])
                for m in METHODS
            ]
        )
    for grid in GRIDS:
        rows.append(
            [grid]
            + [
                _median_max([e["clinical"] for e in disc_errors(records, m, grid=grid)])
                for m in METHODS
            ]
        )
    for shape in SHAPES:
        rows.append(
            [shape]
            + [
                _median_max([e["clinical"] for e in disc_errors(records, m, shape=shape)])
                for m in METHODS
            ]
        )
    lines += table(["Subset", *METHODS], rows)

    lines += [
        "### Time per structure",
        "",
        "Median seconds for one structure and one dose, from nothing.",
        "",
    ]
    rows = []
    for radius in RADII:
        row = [f"R = {radius:g} mm"]
        for m in METHODS:
            times = [
                r["result"].seconds for r in records if r["method"] == m and r["radius"] == radius
            ]
            row.append(f"{np.median(times):.2f}")
        rows.append(row)
    lines += table(["Size", *METHODS], rows)
    return lines


def failures(records: list[dict]) -> list[str]:
    seen: dict[tuple[str, str], int] = defaultdict(int)
    for r in records:
        if r["result"].error:
            error = r["result"].error
            seen[(r["method"], error if len(error) <= 160 else error[:157] + "...")] += 1
    outside = defaultdict(int)
    for r in records:
        if r["result"].outside:
            outside[r["method"]] += 1
    lines = []
    if seen:
        lines += table(
            ["Method", "Error", "Cases"], [[m, e, str(n)] for (m, e), n in sorted(seen.items())]
        )
    else:
        lines += ["No method failed on any case.", ""]
    if outside:
        lines += [
            "Cases with samples outside the dose grid: "
            + ", ".join(f"{m} {n}" for m, n in sorted(outside.items())),
            "",
        ]
    return lines


def large_section(records: list[dict]) -> list[str]:
    names = [variant_name(m, s) for m, s in large_variants()]
    found = {(r["shape"], r["radius"], r["method"]): r for r in records}
    sizes = list(dict.fromkeys((r["shape"], r["radius"]) for r in records))

    def label(shape: str, radius: float) -> str:
        volume = found[(shape, radius, names[0])]["truth"]["volume_cc"]
        return f"{shape}, R {radius:g} mm ({volume:,.0f} cc)"

    timing, accuracy = [], []
    for shape, radius in sizes:
        t_row, a_row = [label(shape, radius)], [label(shape, radius)]
        for name in names:
            r = found[(shape, radius, name)]
            result, truth = r["result"], r["truth"]
            if result.error:
                t_row.append("failed")
                a_row.append("failed")
                continue
            samples = f" ({result.samples / 1e6:.1f} M)" if result.samples else ""
            t_row.append(f"{result.seconds:.2f}{samples}")
            worst = max(abs(result.metrics[k] - truth[k]) for k in CLINICAL) / LARGE_SLOPE
            volume = pct(result.metrics["volume_cc"], truth["volume_cc"])
            a_row.append(f"{worst:.2f} mm / {volume:+.2f} %")
        timing.append(t_row)
        accuracy.append(a_row)
    lines = [
        "## Large structures: time against accuracy",
        "",
        f"One sphere or cylinder per size (the cylinders as tall as they are wide) on the "
        f"tender cohort's grid: CT 0.98 × 0.98 × 2 mm, dose 2 mm. The dose is "
        f"{LARGE_D0:g} Gy at the centre rising {LARGE_SLOPE:g} Gy/mm obliquely, gentler "
        "than elsewhere so it stays positive across 20 cm.",
        "",
        "### Seconds for one structure against one dose",
        "",
        "From reading the contours to the statistics, single-threaded, on the machine "
        f"named under Environment; a run under {LARGE_RETIME_S:g} s is the faster of two. "
        "In brackets: dose samples taken, in millions.",
        "",
    ]
    lines += table(["Structure", *names], timing)
    lines += [
        "### Accuracy",
        "",
        "The worst error among Dmean, D99, D95, D5, D1 and D0.03cc, as the boundary shift "
        f"that would cause it (Gy ÷ {LARGE_SLOPE:g}), then the volume error.",
        "",
    ]
    return lines + table(["Structure", *names], accuracy)


def with_corrected_lookup(records: list[dict]) -> list[dict]:
    """dicompyler's records again, with D{x} read from its histogram by our lookup."""
    out = []
    for r in records:
        corrected = r["result"].corrected
        if r["method"] in ("dicompyler", "dicompyler-ss") and corrected:
            result = replace(r["result"], metrics={**r["result"].metrics, **corrected})
            out.append({**r, "method": f"{r['method']} (lookup corrected)", "result": result})
    return out


def lookup_section(nelms: list[dict], discs: list[dict]) -> list[str]:
    extra_nelms, extra_discs = with_corrected_lookup(nelms), with_corrected_lookup(discs)
    names = [
        "dicompyler",
        "dicompyler (lookup corrected)",
        "dicompyler-ss",
        "dicompyler-ss (lookup corrected)",
        "mask-ss",
        "polygon",
    ]
    rows = []
    subset = ("volume_cc", *CLINICAL)
    for test in ("1", "2", "2s") if nelms else ():
        row = [f"Nelms {TEST_NAMES[test].split(':')[0]}: beyond 3 %, without Dmin, Dmax"]
        for name in names:
            c = nelms_counts(nelms + extra_nelms, test, name)
            row.append(f"{sum(c[k][0] for k in subset)}/{sum(c[k][1] for k in subset)}")
        rows.append(row)
    if discs:
        row = ["Disc phantoms: worst of Dmean-D0.03cc, median / largest (Gy)"]
        for name in names:
            row.append(_median_max([e["clinical"] for e in disc_errors(discs + extra_discs, name)]))
        rows.append(row)
    lines = [
        "## dicompyler with its lookup corrected",
        "",
        "dicompyler's own histogram, with D*x* read the way the other methods read it (the "
        "lowest dose the hottest *x* receives, to the 1 cGy bin). The volume, Dmin, Dmax "
        "and Dmean are unchanged. This separates what dicompyler's lookup costs (next "
        "section) from what its sampling costs, and is what keeping dicompyler with only "
        "the lookup replaced would give.",
        "",
    ]
    return lines + table(["", *names], rows)


def zero_doses(records: list[dict]) -> list[str]:
    """Dose statistics reported as exactly 0 Gy where the truth is well above it."""
    counts: dict[tuple[str, str], dict[str, int]] = defaultdict(lambda: defaultdict(int))
    cases: dict[tuple[str, str], int] = defaultdict(int)
    for r in records:
        if r["method"] not in METHODS:
            continue
        truth = r["row"].truth if r["benchmark"] == "nelms" else r["truth"]
        key = (r["benchmark"], r["method"])
        hit = False
        for metric in METRICS[1:]:
            if r["result"].metrics[metric] == 0.0 and truth[metric] > 0.5:
                counts[key][metric] += 1
                hit = True
        cases[key] += hit
    rows = []
    for (benchmark, method), per_metric in sorted(counts.items()):
        total = sum(1 for r in records if r["benchmark"] == benchmark and r["method"] == method)
        rows.append(
            [
                "Nelms" if benchmark == "nelms" else "Disc phantoms",
                method,
                f"{cases[(benchmark, method)]}/{total}",
                ", ".join(f"{LABEL[m]} {n}" for m, n in per_metric.items()),
            ]
        )
    if not rows:
        return ["No method reported a dose of 0 Gy where the truth is above it.", ""]
    lines = [
        "dicompyler-core's `dose_constraint` answers D*x* with the first dose bin whose "
        "cumulative volume is nearest *x*. Every bin from 0 Gy up to Dmin holds 100 % of "
        "the volume, so when no later bin is nearer *x* than 100 % is, the answer is the "
        "first bin: 0 Gy. For D99 that happens once the coldest 1 cGy bin holds more "
        "than 2 % of the volume, for D95 more than 10 %; a structure whose dose falls in "
        "a single bin gets 0 Gy for every D*x*, and a D*x*cc larger than the volume "
        "dicompyler found gets 0 Gy too. Few dose samples, or whole planes sharing one "
        "dose, are enough.",
        "",
    ]
    return lines + table(["Benchmark", "Method", "Cases", "Statistics at 0 Gy"], rows, align="llrl")


def _version(package: str) -> str:
    try:
        return importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:
        return "not installed"


def _commit() -> str:
    try:
        out = subprocess.run(
            ["git", "rev-parse", "--short", "HEAD"],
            capture_output=True,
            text=True,
            cwd=Path(__file__).parent,
            check=True,
        )
        return out.stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def write_report(
    target: Path,
    nelms: list[dict],
    discs: list[dict],
    large: list[dict],
    checks: dict[str, str],
    args: argparse.Namespace,
) -> None:
    lines = [
        "# DVH method validation",
        "",
        f"Generated {datetime.date.today().isoformat()} by `scripts/validate_dvh_methods.py` "
        f"(AutoSeg {autoseg_version}, commit {_commit()}). Regenerate with "
        "`python scripts/validate_dvh_methods.py --nelms <folder> --out "
        "docs/DVH_METHOD_VALIDATION.md`.",
        "",
        "Roadmap item #7: should structure-set DVHs keep coming from dicompyler-core, or "
        "come from the same contour reading as the geometric metrics? Every candidate "
        "is scored here against DVHs whose true values are known exactly. This report "
        "records the measurements; the decision is recorded separately.",
        "",
        "## Methods",
        "",
    ]
    lines += table(
        ["Method", "Contours read by", "Dose sampled at", "Each sample stands for"],
        [
            [
                "dicompyler",
                "dicompyler-core: every loop tested, loops combined by exclusive-or",
                "each dose-grid point in each contour plane; the dose interpolated "
                "between dose planes only",
                "one dose voxel's area × the gap between contour planes",
            ],
            [
                "dicompyler-ss",
                "the same",
                "a grid a quarter of the dose pixel, in each contour plane",
                "that grid's cell × the gap between contour planes",
            ],
            [
                "mask",
                "the shared reading, half-open fill (the 3D metrics' own mask)",
                "each CT voxel centre, trilinear",
                "one CT voxel",
            ],
            [
                "mask-ss",
                "the same mask",
                f"sub-samples ≤ {args.subsample_mm:g} mm apart in every CT voxel, trilinear",
                "an equal share of its voxel",
            ],
            [
                "polygon",
                "the shared reading's regions, no voxels",
                f"sub-cells ≤ {args.subsample_mm:g} mm apart, trilinear, each at the "
                "centroid of the part covered",
                "the exact area of the region in its sub-cell × its share of the slice",
            ],
        ],
        align="llll",
    )
    lines += [
        "Every method treats a contour as a slab one slice thick, centred on its plane, "
        "and none interpolates between contours: that is the convention in both "
        "benchmarks' truth. The sub-sampling spacing is rounded to an odd number of "
        "sub-samples per voxel edge, so one always sits on the voxel centre. D*x* is the "
        "lowest dose the hottest *x* of the volume receives. The three sampled methods "
        "take a slice at a time into a dose histogram of 1 mGy bins, so memory does not "
        "grow with the number of samples; Dmin, Dmax and Dmean are kept exactly and D*x* "
        "is read to the bin's centre. Polygon finds each sub-cell's covered area and its "
        "centroid exactly, without clipping, by the signed-area accumulation fonts are "
        "rasterised with, so its extra work grows with a structure's outline rather than "
        "its area. The dicompyler rows are "
        "production's numbers: each case is checked against "
        "`core.dvh.compute_dvh_metrics`.",
        "",
        "## Nelms et al. 2015",
        "",
        "Nelms B, Stambaugh C, Hunt D, Tonner B, Zhang G, Feygelman V. *Methods, software "
        "and datasets to verify DVH calculations against analytical values: twenty years "
        "late(r).* Med Phys 2015;42(8):4435-48. doi:10.1118/1.4923175. PMID 26233174.",
        "",
        "A sphere, axial and rotated cylinders, and axial and rotated cones, 24 mm across, "
        "contoured every 0.2, 1, 2 or 3 mm on a CT of 0.6 mm pixels, in 1 Gy/mm linear dose "
        "fields (16 Gy at the centre) along anterior-posterior or superior-inferior. The "
        "truth is the solid itself extended half a slice beyond its end contours, so a "
        "method is also charged for filling the gap between contour planes: a rotated "
        "cylinder is a stack of rectangles, and no slab method can recover its round "
        "side. Pinnacle3 v9.8 and PlanIQ v2.1 are the paper's own results on the same "
        "data. The datasets are the paper's supplementary material and are not "
        "redistributed with AutoSeg.",
        "",
    ]
    if nelms:
        for test in ("1", "2", "2s"):
            lines += nelms_section(nelms, test)
        lines += ["### The same at 2 %", ""] + nelms_two_percent(nelms)
        lines += nelms_test3(nelms)
    else:
        lines += ["Not run: no Nelms data were given.", ""]

    if discs:
        lines += discs_section(discs, args.placements)
    if nelms:
        lines += sweep_section(nelms, args.subsample_mm)
    if large:
        lines += large_section(large)

    lines += lookup_section(nelms, discs)
    lines += ["## Doses reported as 0 Gy", ""] + zero_doses(nelms + discs)
    lines += ["## Failures", ""] + failures(nelms + discs + large)
    lines += [
        "## Checks on the harness",
        "",
        f"- Closed-form disc truth against brute-force integration over the written "
        f"polygons (96-point Gauss-Legendre through each slab, exact polygon clipping), "
        f"largest difference as a share of the structure's volume: {checks['truth']}.",
        f"- Trilinear dose sampling here against SimpleITK's resampling, which the "
        f"consensus DVH uses today, at every voxel of a Nelms sphere: largest difference "
        f"{checks['sampler']}.",
        "- Polygon's sub-cell coverage and centroids, found by accumulation, against "
        f"clipping every sub-cell with shapely (Nelms and disc outlines, 1-5 sub-cells "
        f"per voxel edge): {checks['cells']}.",
        "- The dicompyler rows equal `compute_dvh_metrics` on every case (checked per case; "
        "a mismatch would appear under Failures).",
        "",
        "## Environment",
        "",
        f"{platform.processor() or platform.machine()}, {os.cpu_count()} logical cores; "
        f"Python {platform.python_version()} on {platform.system()} {platform.release()}; "
        f"numpy {_version('numpy')}, scipy {_version('scipy')}, shapely {_version('shapely')}, "
        f"SimpleITK {_version('SimpleITK')}, pydicom {_version('pydicom')}, dicompyler-core "
        f"{_version('dicompyler-core')}. Sub-sample spacing {args.subsample_mm:g} mm; "
        f"{args.placements} placements per disc phantom; seed {args.seed}.",
        "",
    ]
    if nelms:
        lines += ["## Appendix: Test 3 by dataset", "", "Lowest to highest volume error (%)."]
        lines += [""] + nelms_test3_detail(nelms)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("\n".join(lines), encoding="utf-8")


def write_csv(target: Path, nelms: list[dict], discs: list[dict], large: list[dict]) -> None:
    with target.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "benchmark",
                "set",
                "structure",
                "radius_mm",
                "gradient",
                "placement",
                "method",
                "metric",
                "value",
                "truth",
                "difference",
                "percent",
                "seconds",
                "error",
            ]
        )
        for r in nelms:
            row = r["row"]
            for metric in METRICS:
                value, truth = r["result"].metrics[metric], row.truth[metric]
                writer.writerow(
                    [
                        "nelms",
                        row.test,
                        row.structure,
                        "",
                        f"{row.gradient} {row.voxel}",
                        "",
                        r["method"],
                        metric,
                        value,
                        truth,
                        value - truth,
                        pct(value, truth),
                        r["result"].seconds,
                        r["result"].error,
                    ]
                )
        for r in discs:
            for metric in METRICS:
                value, truth = r["result"].metrics[metric], r["truth"][metric]
                writer.writerow(
                    [
                        "discs",
                        r["grid"],
                        r["shape"],
                        r["radius"],
                        r["direction"],
                        r["placement"],
                        r["method"],
                        metric,
                        value,
                        truth,
                        value - truth,
                        pct(value, truth),
                        r["result"].seconds,
                        r["result"].error,
                    ]
                )
        for r in large:
            for metric in METRICS:
                value, truth = r["result"].metrics[metric], r["truth"][metric]
                writer.writerow(
                    [
                        "large",
                        "",
                        r["shape"],
                        r["radius"],
                        "xz",
                        "",
                        r["method"],
                        metric,
                        value,
                        truth,
                        value - truth,
                        pct(value, truth),
                        r["result"].seconds,
                        r["result"].error,
                    ]
                )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--nelms", type=Path, help="Folder of the Nelms et al. 2015 dataset.")
    parser.add_argument("--out", type=Path, required=True, help="Write the markdown report here.")
    parser.add_argument("--csv", type=Path, help="Also write every value, long format.")
    parser.add_argument("--subsample-mm", type=float, default=0.25)
    parser.add_argument("--placements", type=int, default=8, help="Placements per disc phantom.")
    parser.add_argument("--seed", type=int, default=20260924)
    parser.add_argument("--skip-discs", action="store_true")
    parser.add_argument("--skip-large", action="store_true")
    args = parser.parse_args(argv)

    checks = {
        "truth": check_truth(args.seed),
        "sampler": check_sampler(args.nelms),
        "cells": check_polygon_cells(args.nelms, args.seed),
    }
    print(f"checks: {checks}", flush=True)
    nelms: list[dict] = []
    if args.nelms:
        variants: list[tuple[str, float | None]] = [(m, None) for m in METHODS]
        variants += [(m, s) for m in SAMPLED for s in (1.0, 0.5) if s != args.subsample_mm]
        nelms = run_nelms(args.nelms, args.subsample_mm, variants)
    discs = [] if args.skip_discs else run_discs(args.subsample_mm, args.placements, args.seed)
    large = [] if args.skip_large else run_large(args.seed)
    write_report(args.out, nelms, discs, large, checks, args)
    if args.csv:
        write_csv(args.csv, nelms, discs, large)
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
